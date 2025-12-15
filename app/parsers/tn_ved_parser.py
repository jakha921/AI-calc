"""
Парсер кодов ТН ВЭД из Word документов
"""
import re
from typing import List, Dict, Optional, Tuple
from pathlib import Path
import logging

try:
    from docx import Document
except ImportError:
    Document = None

import pandas as pd

logger = logging.getLogger(__name__)


class TNVedParser:
    """Парсер кодов ТН ВЭД из .docx и .xlsx файлов"""

    # Pattern for TN VED code (4-10 digits)
    CODE_PATTERN = re.compile(r"^\d{4,10}$")

    def __init__(self):
        self.codes: List[Dict] = []

    def parse_docx(self, file_path: str) -> List[Dict]:
        """
        Parse TN VED codes from Word document

        Args:
            file_path: Path to .docx file

        Returns:
            List of dictionaries with code data
        """
        if Document is None:
            raise ImportError("python-docx is required. Install with: pip install python-docx")

        doc = Document(file_path)
        codes = []

        # Try to parse from tables first
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 2:
                    code_data = self._parse_row(cells)
                    if code_data:
                        codes.append(code_data)

        # If no tables, try to parse from paragraphs
        if not codes:
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    code_data = self._parse_text_line(text)
                    if code_data:
                        codes.append(code_data)

        self.codes = codes
        logger.info(f"Parsed {len(codes)} TN VED codes from {file_path}")
        return codes

    def parse_excel(self, file_path: str, code_col: int = 0, desc_col: int = 1) -> List[Dict]:
        """
        Parse TN VED codes from Excel file

        Args:
            file_path: Path to .xlsx file
            code_col: Column index for codes
            desc_col: Column index for descriptions

        Returns:
            List of dictionaries with code data
        """
        df = pd.read_excel(file_path, dtype=str)
        codes = []

        for _, row in df.iterrows():
            try:
                code = str(row.iloc[code_col]).strip()
                description = str(row.iloc[desc_col]).strip() if desc_col < len(row) else ""

                # Clean code
                code = re.sub(r"\D", "", code)

                if self._validate_code(code):
                    codes.append({
                        "code": code.zfill(10) if len(code) < 10 else code[:10],
                        "description": description,
                        "level": self._get_level(code),
                        "parent_code": self._get_parent_code(code),
                    })
            except Exception as e:
                logger.warning(f"Error parsing row: {e}")
                continue

        self.codes = codes
        logger.info(f"Parsed {len(codes)} TN VED codes from {file_path}")
        return codes

    def parse_csv(self, file_path: str, code_col: str = "code", desc_col: str = "description") -> List[Dict]:
        """Parse TN VED codes from CSV file"""
        df = pd.read_csv(file_path, dtype=str)
        codes = []

        for _, row in df.iterrows():
            code = str(row.get(code_col, "")).strip()
            description = str(row.get(desc_col, "")).strip()

            code = re.sub(r"\D", "", code)

            if self._validate_code(code):
                codes.append({
                    "code": code.zfill(10) if len(code) < 10 else code[:10],
                    "description": description,
                    "level": self._get_level(code),
                    "parent_code": self._get_parent_code(code),
                })

        self.codes = codes
        logger.info(f"Parsed {len(codes)} TN VED codes from {file_path}")
        return codes

    def _parse_row(self, cells: List[str]) -> Optional[Dict]:
        """Parse a table row"""
        for i, cell in enumerate(cells):
            # Try to find code in cell
            clean_cell = re.sub(r"\s+", "", cell)
            code_match = re.search(r"\d{4,10}", clean_cell)

            if code_match:
                code = code_match.group()
                # Get description from next cell or rest of current cell
                if i + 1 < len(cells):
                    description = cells[i + 1]
                else:
                    description = cell.replace(code, "").strip(" -–—")

                if self._validate_code(code):
                    return {
                        "code": code.zfill(10) if len(code) < 10 else code[:10],
                        "description": description,
                        "level": self._get_level(code),
                        "parent_code": self._get_parent_code(code),
                    }
        return None

    def _parse_text_line(self, text: str) -> Optional[Dict]:
        """Parse a text line like '8703231981 - Автомобили легковые'"""
        # Pattern: CODE separator DESCRIPTION
        match = re.match(r"(\d{4,10})\s*[-–—:]\s*(.+)", text)
        if match:
            code = match.group(1)
            description = match.group(2).strip()
            if self._validate_code(code):
                return {
                    "code": code.zfill(10) if len(code) < 10 else code[:10],
                    "description": description,
                    "level": self._get_level(code),
                    "parent_code": self._get_parent_code(code),
                }
        return None

    def _validate_code(self, code: str) -> bool:
        """Validate TN VED code format"""
        if not code:
            return False
        clean_code = re.sub(r"\D", "", code)
        return 4 <= len(clean_code) <= 10

    def _get_level(self, code: str) -> int:
        """
        Get hierarchy level based on code length:
        - 2 digits: Group (Группа)
        - 4 digits: Position (Позиция)
        - 6 digits: Subposition (Субпозиция)
        - 8-10 digits: Item (Подсубпозиция)
        """
        code = code.rstrip("0") if code else ""
        length = len(code)
        if length <= 2:
            return 2
        elif length <= 4:
            return 4
        elif length <= 6:
            return 6
        elif length <= 8:
            return 8
        return 10

    def _get_parent_code(self, code: str) -> Optional[str]:
        """Get parent code for hierarchy"""
        if len(code) <= 4:
            return None
        elif len(code) <= 6:
            return code[:4].ljust(10, "0")
        elif len(code) <= 8:
            return code[:6].ljust(10, "0")
        else:
            return code[:8].ljust(10, "0")

    def get_hierarchy(self) -> Dict[str, List[Dict]]:
        """Organize codes into hierarchy"""
        hierarchy = {
            "groups": [],  # 2-digit
            "positions": [],  # 4-digit
            "subpositions": [],  # 6-digit
            "items": [],  # 8-10 digit
        }

        for code_data in self.codes:
            level = code_data["level"]
            if level == 2:
                hierarchy["groups"].append(code_data)
            elif level == 4:
                hierarchy["positions"].append(code_data)
            elif level == 6:
                hierarchy["subpositions"].append(code_data)
            else:
                hierarchy["items"].append(code_data)

        return hierarchy
